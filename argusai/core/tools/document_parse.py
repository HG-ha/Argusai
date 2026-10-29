"""Document parsing — local (pypdf / python-docx) + optional MinerU cloud."""

from __future__ import annotations

import io
import time
import zipfile
from pathlib import Path
from typing import Any

from .common import _session_sandbox_root

_DOC_EXTS = {".pdf", ".docx"}
_MINERU_EXTS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".jp2", ".webp", ".gif", ".bmp",
    ".doc", ".docx", ".ppt", ".pptx", ".xls", ".xlsx",
}
_DEFAULT_MINERU_BASE = "https://mineru.net"
_MAX_OUTPUT_CHARS = 200_000
_MAX_FILE_BYTES = 32 * 1024 * 1024


def parse_document(
    api: Any,
    path: str = "",
    engine: str = "auto",
    pages: str = "",
    output_path: str = "",
    language: str = "ch",
    ocr: bool = False,
    max_chars: int = 0,
    **kwargs: Any,
) -> dict[str, Any]:
    """Parse a sandbox document into markdown/text.

    **Local (default):** ``.pdf`` via ``pypdf``, ``.docx`` via ``python-docx``
    — no network, small deps. Best for digital (non-scan) office files.

    **MinerU:** when ``engine='mineru'`` or ``auto`` cannot handle the file
    locally and MinerU is enabled in Settings → Automation. Uploads the file
    to mineru.net (token optional: Agent light API without token; precise API
    when an API key is configured).

    ``path`` must stay inside the session sandbox (e.g. ``_attachments/a.pdf``).
    """
    root = _session_sandbox_root(api, kwargs)
    if not str(path or "").strip():
        return {"ok": False, "error": "path is required"}

    from argusai.core.workspace.session_sandbox import resolve_in_sandbox

    try:
        target = resolve_in_sandbox(root, path, allowed_exts=_DOC_EXTS | _MINERU_EXTS)
    except ValueError as exc:
        return {"ok": False, "error": str(exc)}
    if not target.is_file():
        return {"ok": False, "error": f"file not found: {path}"}

    size = target.stat().st_size
    if size <= 0:
        return {"ok": False, "error": "file is empty"}
    if size > _MAX_FILE_BYTES:
        return {
            "ok": False,
            "error": f"file too large ({size} bytes > {_MAX_FILE_BYTES})",
        }

    ext = target.suffix.lower()
    engine_key = str(engine or "auto").strip().lower() or "auto"
    settings = _document_parse_settings(api)
    cap = int(max_chars) if int(max_chars or 0) > 0 else int(settings.get("max_output_chars") or _MAX_OUTPUT_CHARS)
    cap = max(1000, min(cap, 2_000_000))

    local_ok = ext in _DOC_EXTS and not bool(ocr)
    use_mineru = False
    if engine_key == "mineru":
        use_mineru = True
    elif engine_key == "local":
        use_mineru = False
    else:  # auto
        use_mineru = (not local_ok) and bool(settings.get("mineru_enabled"))

    if use_mineru:
        if not settings.get("mineru_enabled"):
            return {
                "ok": False,
                "error": "MinerU is disabled. Enable it in Settings → Automation → MinerU, or use engine='local' for pdf/docx.",
            }
        if ext not in _MINERU_EXTS:
            return {"ok": False, "error": f"MinerU does not support extension {ext}"}
        result = _parse_via_mineru(
            target,
            settings=settings,
            pages=str(pages or "").strip(),
            language=str(language or "ch").strip() or "ch",
            ocr=bool(ocr),
        )
    else:
        if not local_ok:
            if settings.get("mineru_enabled"):
                return {
                    "ok": False,
                    "error": (
                        f"local engine cannot parse {ext}"
                        + (" with ocr=true" if ocr else "")
                        + "; retry with engine='mineru'"
                    ),
                }
            return {
                "ok": False,
                "error": (
                    f"local engine supports only .pdf/.docx (got {ext}). "
                    "Enable MinerU in Settings for broader formats / OCR."
                ),
            }
        try:
            if ext == ".pdf":
                text, meta = _parse_pdf_local(target, pages=str(pages or "").strip())
            else:
                text, meta = _parse_docx_local(target)
        except Exception as exc:
            return {"ok": False, "error": f"local parse failed: {exc}"[:300]}
        result = {
            "ok": True,
            "engine": "local",
            "provider": meta.get("provider") or ext.lstrip("."),
            "markdown": text,
            "meta": meta,
        }

    if not result.get("ok"):
        return result

    markdown = str(result.get("markdown") or "")
    truncated = False
    if len(markdown) > cap:
        markdown = markdown[:cap] + f"\n\n…[truncated at {cap} chars]"
        truncated = True

    out_rel = ""
    if str(output_path or "").strip():
        from argusai.core.workspace.session_sandbox import write_file

        try:
            written = write_file(root, str(output_path).strip(), markdown, file_type="md")
            out_rel = str(written.get("path") or output_path)
        except Exception as exc:
            return {
                "ok": False,
                "error": f"parsed ok but failed to write output_path: {exc}"[:300],
                "markdown": markdown[:2000],
                "engine": result.get("engine"),
            }
    else:
        # Default spill for long results so the model can re-read via sandbox tools.
        if len(markdown) > 12_000:
            from argusai.core.workspace.session_sandbox import write_file

            stem = target.stem or "document"
            auto_out = f"_attachments/{stem}.parsed.md"
            try:
                written = write_file(root, auto_out, markdown, file_type="md")
                out_rel = str(written.get("path") or auto_out)
            except Exception:
                out_rel = ""

    return {
        "ok": True,
        "engine": result.get("engine"),
        "provider": result.get("provider"),
        "path": path,
        "output_path": out_rel,
        "chars": len(markdown),
        "truncated": truncated,
        "markdown": markdown if len(markdown) <= 12_000 or not out_rel else markdown[:4000] + f"\n\n…[full text at {out_rel}]",
        "meta": result.get("meta") or {},
    }


def _document_parse_settings(api: Any) -> dict[str, Any]:
    try:
        settings = api.settings_mgr.get()
    except Exception:
        return {
            "mineru_enabled": False,
            "mineru_api_key": "",
            "mineru_base_url": _DEFAULT_MINERU_BASE,
            "mineru_model_version": "vlm",
            "mineru_timeout_seconds": 300,
            "max_output_chars": _MAX_OUTPUT_CHARS,
        }
    return {
        "mineru_enabled": bool(getattr(settings, "document_parse_mineru_enabled", False)),
        "mineru_api_key": str(getattr(settings, "document_parse_mineru_api_key", "") or "").strip(),
        "mineru_base_url": str(
            getattr(settings, "document_parse_mineru_base_url", "") or _DEFAULT_MINERU_BASE
        ).rstrip("/"),
        "mineru_model_version": str(
            getattr(settings, "document_parse_mineru_model_version", "") or "vlm"
        ).strip()
        or "vlm",
        "mineru_timeout_seconds": float(
            getattr(settings, "document_parse_mineru_timeout_seconds", 300) or 300
        ),
        "max_output_chars": int(
            getattr(settings, "document_parse_max_output_chars", _MAX_OUTPUT_CHARS)
            or _MAX_OUTPUT_CHARS
        ),
    }


def _parse_pdf_local(path: Path, *, pages: str = "") -> tuple[str, dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    total = len(reader.pages)
    indexes = _page_indexes(pages, total)
    parts: list[str] = []
    for i in indexes:
        text = reader.pages[i].extract_text() or ""
        parts.append(f"## Page {i + 1}\n\n{text.strip()}".rstrip())
    body = "\n\n".join(p for p in parts if p.strip())
    if not body.strip():
        raise ValueError(
            "no extractable text (likely a scanned PDF). "
            "Enable MinerU and retry with engine='mineru' ocr=true"
        )
    return body, {
        "provider": "pypdf",
        "page_count": total,
        "pages_parsed": [i + 1 for i in indexes],
    }


def _parse_docx_local(path: Path) -> tuple[str, dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    blocks: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").rstrip()
        if text:
            blocks.append(text)
    for table in doc.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [" ".join((c.text or "").split()) for c in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            header_sep = "| " + " | ".join("---" for _ in row.cells) + " |"
            blocks.append(rows[0])
            blocks.append(header_sep)
            blocks.extend(rows[1:])
    body = "\n\n".join(blocks).strip()
    if not body:
        raise ValueError("docx contained no extractable paragraphs/tables")
    return body, {"provider": "python-docx", "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}


def _page_indexes(pages: str, total: int) -> list[int]:
    if total <= 0:
        return []
    raw = str(pages or "").strip()
    if not raw:
        return list(range(total))
    if "-" in raw and "," not in raw:
        left, _, right = raw.partition("-")
        start = max(1, int(left or "1"))
        end = min(total, int(right or str(total)))
        if end < start:
            start, end = end, start
        return list(range(start - 1, end))
    if raw.isdigit():
        n = int(raw)
        if 1 <= n <= total:
            return [n - 1]
        raise ValueError(f"page out of range: {n} (doc has {total} pages)")
    raise ValueError("pages must be a single page (5) or a range (1-10)")


def _parse_via_mineru(
    path: Path,
    *,
    settings: dict[str, Any],
    pages: str,
    language: str,
    ocr: bool,
) -> dict[str, Any]:
    from argusai.core.infra.httpx_client import direct_async_client

    base = str(settings.get("mineru_base_url") or _DEFAULT_MINERU_BASE).rstrip("/")
    token = str(settings.get("mineru_api_key") or "").strip()
    timeout = max(30.0, min(float(settings.get("mineru_timeout_seconds") or 300), 1800.0))
    model_version = str(settings.get("mineru_model_version") or "vlm").strip() or "vlm"
    data = path.read_bytes()

    async def _run() -> dict[str, Any]:
        async with direct_async_client(timeout=timeout, follow_redirects=True) as client:
            if token:
                return await _mineru_precise_upload(
                    client,
                    base=base,
                    token=token,
                    filename=path.name,
                    data=data,
                    model_version=model_version,
                    timeout=timeout,
                )
            return await _mineru_agent_upload(
                client,
                base=base,
                filename=path.name,
                data=data,
                pages=pages,
                language=language,
                ocr=ocr,
                timeout=timeout,
            )

    try:
        return _run_blocking(_run)
    except Exception as exc:
        return {"ok": False, "error": f"MinerU request failed: {exc}"[:400]}


def _run_blocking(coro_factory: Any) -> Any:
    import asyncio
    import concurrent.futures

    try:
        asyncio.get_running_loop()
        nested = True
    except RuntimeError:
        nested = False
    if nested:
        with concurrent.futures.ThreadPoolExecutor() as pool:
            return pool.submit(asyncio.run, coro_factory()).result()
    return asyncio.run(coro_factory())


async def _mineru_agent_upload(
    client: Any,
    *,
    base: str,
    filename: str,
    data: bytes,
    pages: str,
    language: str,
    ocr: bool,
    timeout: float,
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "file_name": filename,
        "language": language or "ch",
        "enable_table": True,
        "is_ocr": bool(ocr),
        "enable_formula": True,
    }
    if pages:
        payload["page_range"] = pages
    resp = await client.post(f"{base}/api/v1/agent/parse/file", json=payload)
    body = resp.json()
    if not isinstance(body, dict) or int(body.get("code", -1)) != 0:
        return {"ok": False, "error": f"MinerU agent submit failed: {(body.get('msg') if isinstance(body, dict) else None) or body}"[:300]}
    info = body.get("data") if isinstance(body.get("data"), dict) else {}
    task_id = str(info.get("task_id") or "")
    file_url = str(info.get("file_url") or "")
    if not task_id or not file_url:
        return {"ok": False, "error": "MinerU agent submit missing task_id/file_url"}
    put = await client.put(file_url, content=data)
    if put.status_code >= 300:
        return {"ok": False, "error": f"MinerU upload failed HTTP {put.status_code}"}
    markdown = await _poll_mineru_agent(client, base=base, task_id=task_id, timeout=timeout)
    return {
        "ok": True,
        "engine": "mineru",
        "provider": "mineru-agent",
        "markdown": markdown,
        "meta": {"task_id": task_id, "api": "agent"},
    }


async def _poll_mineru_agent(client: Any, *, base: str, task_id: str, timeout: float) -> str:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        resp = await client.get(f"{base}/api/v1/agent/parse/{task_id}")
        body = resp.json()
        data = body.get("data") if isinstance(body.get("data"), dict) else {}
        state = str(data.get("state") or "")
        if state == "done":
            md_url = str(data.get("markdown_url") or "")
            if not md_url:
                raise RuntimeError("MinerU done but markdown_url missing")
            md_resp = await client.get(md_url)
            md_resp.raise_for_status()
            return md_resp.text
        if state == "failed":
            raise RuntimeError(str(data.get("err_msg") or "MinerU parse failed")[:300])
        await _async_sleep(2.0)
    raise TimeoutError(f"MinerU agent parse timed out after {int(timeout)}s")


async def _mineru_precise_upload(
    client: Any,
    *,
    base: str,
    token: str,
    filename: str,
    data: bytes,
    model_version: str,
    timeout: float,
) -> dict[str, Any]:
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    resp = await client.post(
        f"{base}/api/v4/file-urls/batch",
        headers=headers,
        json={
            "files": [{"name": filename, "data_id": filename}],
            "model_version": model_version,
        },
    )
    body = resp.json()
    if not isinstance(body, dict) or int(body.get("code", -1)) != 0:
        return {"ok": False, "error": f"MinerU precise submit failed: {(body.get('msg') if isinstance(body, dict) else None) or body}"[:300]}
    info = body.get("data") if isinstance(body.get("data"), dict) else {}
    batch_id = str(info.get("batch_id") or "")
    urls = list(info.get("file_urls") or [])
    if not batch_id or not urls:
        return {"ok": False, "error": "MinerU precise submit missing batch_id/file_urls"}
    put = await client.put(str(urls[0]), content=data)
    if put.status_code >= 300:
        return {"ok": False, "error": f"MinerU upload failed HTTP {put.status_code}"}
    markdown = await _poll_mineru_precise(client, base=base, token=token, batch_id=batch_id, timeout=timeout)
    return {
        "ok": True,
        "engine": "mineru",
        "provider": "mineru-precise",
        "markdown": markdown,
        "meta": {"batch_id": batch_id, "api": "precise", "model_version": model_version},
    }


async def _poll_mineru_precise(
    client: Any,
    *,
    base: str,
    token: str,
    batch_id: str,
    timeout: float,
) -> str:
    headers = {"Authorization": f"Bearer {token}"}
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        resp = await client.get(
            f"{base}/api/v4/extract-results/batch/{batch_id}",
            headers=headers,
        )
        body = resp.json()
        data = body.get("data") if isinstance(body.get("data"), dict) else {}
        results = data.get("extract_result")
        if isinstance(results, dict):
            results = [results]
        if not isinstance(results, list) or not results:
            await _async_sleep(2.0)
            continue
        item = results[0] if isinstance(results[0], dict) else {}
        state = str(item.get("state") or "")
        if state == "done":
            zip_url = str(item.get("full_zip_url") or "")
            if not zip_url:
                raise RuntimeError("MinerU done but full_zip_url missing")
            zip_resp = await client.get(zip_url)
            zip_resp.raise_for_status()
            return _markdown_from_mineru_zip(zip_resp.content)
        if state == "failed":
            raise RuntimeError(str(item.get("err_msg") or "MinerU parse failed")[:300])
        await _async_sleep(2.0)
    raise TimeoutError(f"MinerU precise parse timed out after {int(timeout)}s")


def _markdown_from_mineru_zip(blob: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        names = zf.namelist()
        preferred = [n for n in names if n.endswith("full.md") or n.endswith("/full.md") or n == "full.md"]
        pick = preferred[0] if preferred else next((n for n in names if n.endswith(".md")), "")
        if not pick:
            raise RuntimeError("MinerU zip contained no markdown")
        return zf.read(pick).decode("utf-8", errors="replace")


async def _async_sleep(seconds: float) -> None:
    import asyncio

    await asyncio.sleep(seconds)
